"""Turn an uploaded file into plain text.

Only formats that can be read faithfully are accepted. A scanned PDF has no
text layer, and silently indexing nothing would look like a working upload that
then "knows" nothing - so that case raises with an explanation instead.

Content is checked against the file's magic bytes, not just its extension: a
`.pdf` that is really an executable is rejected before any parser sees it.
"""

import csv
import io
import json

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json", ".yaml", ".yml",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rs", ".c", ".cpp", ".h",
    ".html", ".htm", ".xml", ".sql", ".sh", ".toml", ".ini", ".cfg",
}
BINARY_EXTS = {".pdf", ".docx", ".pptx", ".xlsx"}
SUPPORTED = sorted(TEXT_EXTS | BINARY_EXTS)

_ZIP_MAGIC = b"PK\x03\x04"


class ParseError(ValueError):
    """The file cannot be turned into usable text. Message is safe to show."""


def _ext(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ParseError("this PDF is password-protected")
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[page {i}]\n{text}")
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if texts:
            slides.append(f"[slide {i}]\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        if rows:
            out.append(f"[sheet {ws.title}]\n" + "\n".join(rows))
    return "\n\n".join(out)


def _text(data: bytes, ext: str) -> str:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    if "\x00" in text:
        raise ParseError("this looks like a binary file, not text")
    if ext in (".csv", ".tsv"):
        rows = csv.reader(io.StringIO(text), delimiter="\t" if ext == ".tsv" else ",")
        return "\n".join(" | ".join(r) for r in rows if any(c.strip() for c in r))
    if ext == ".json":
        try:
            return json.dumps(json.loads(text), indent=2)
        except ValueError:
            return text
    return text


def parse_document(filename: str, data: bytes) -> str:
    ext = _ext(filename)
    if ext not in TEXT_EXTS and ext not in BINARY_EXTS:
        raise ParseError(f"unsupported file type {ext or '(none)'}; use one of: "
                         + ", ".join(SUPPORTED))
    if not data:
        raise ParseError("the file is empty")
    if ext == ".pdf" and not data.startswith(b"%PDF"):
        raise ParseError("the file is not a real PDF")
    if ext in (".docx", ".pptx", ".xlsx") and not data.startswith(_ZIP_MAGIC):
        raise ParseError(f"the file is not a real {ext} document")

    try:
        if ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext == ".pptx":
            text = _pptx(data)
        elif ext == ".xlsx":
            text = _xlsx(data)
        else:
            text = _text(data, ext)
    except ParseError:
        raise
    except Exception as e:
        raise ParseError(f"could not read the file ({type(e).__name__})") from e

    text = text.strip()
    if not text:
        hint = " - it may be a scanned image with no text layer" if ext == ".pdf" else ""
        raise ParseError(f"no text could be extracted{hint}")
    return text
